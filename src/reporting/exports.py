"""Exports des tables finales pour Power BI (CSV UTF-8) et rapport Excel automatisé."""
from __future__ import annotations

import logging
from datetime import date

import pandas as pd
from openpyxl import load_workbook
from openpyxl.formatting.rule import ColorScaleRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from src import config, load

log = logging.getLogger(__name__)

TABLES_POWERBI = [
    "dim_commune", "dim_typologie", "dim_date", "fct_hebergements_classes",
    "mart_indicateurs_communes", "mart_qualite_anomalies", "mart_qualite_synthese",
    "mart_suivi_ingestions",
]

BLEU = "1F4E79"


def exporter_csv() -> None:
    dossier = config.EXPORT_DIR / "powerbi"
    dossier.mkdir(parents=True, exist_ok=True)
    for table in TABLES_POWERBI:
        df = load.lire(f"select * from marts.{table}")
        df.to_csv(dossier / f"{table}.csv", index=False, encoding="utf-8")
    log.info("%d tables exportées pour Power BI dans %s", len(TABLES_POWERBI), dossier)


def _mettre_en_forme(ws, largeur_max: int = 60) -> None:
    for cellule in ws[1]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = PatternFill("solid", fgColor=BLEU)
        cellule.alignment = Alignment(vertical="center", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for i, colonne in enumerate(ws.columns, start=1):
        longueur = max(len(str(c.value)) if c.value is not None else 0 for c in colonne)
        ws.column_dimensions[get_column_letter(i)].width = min(max(longueur + 2, 10), largeur_max)


def rapport_excel(synthese: str) -> str:
    chemin = config.EXPORT_DIR / f"Rapport_offre_classee_{config.DEPARTEMENT}_{date.today():%Y-%m}.xlsx"
    communes = load.lire("""
        select nom_commune as "Commune", nom_epci as "Intercommunalité", population as "Population",
               nb_etablissements as "Établissements", capacite_totale as "Capacité (pers.)",
               capacite_hotellerie as "dont hôtellerie", capacite_plein_air as "dont plein air",
               capacite_residences as "dont résidences/villages",
               taux_fonction_touristique as "Taux de fonction touristique",
               nb_classements_a_renouveler as "Classements à renouveler"
        from marts.mart_indicateurs_communes
        where nb_etablissements > 0
        order by capacite_totale desc
    """)
    epci = load.lire("""
        select nom_epci as "Intercommunalité", sum(nb_etablissements) as "Établissements",
               sum(capacite_totale) as "Capacité (pers.)", sum(population) as "Population",
               round(100.0 * sum(capacite_totale) / nullif(sum(population), 0), 1) as "Taux de fonction touristique"
        from marts.mart_indicateurs_communes group by 1 having sum(nb_etablissements) > 0
        order by 3 desc
    """)
    a_renouveler = load.lire("""
        select nom_commercial as "Établissement", typologie as "Typologie", commune_source as "Commune",
               classement as "Classement", date_classement as "Classé le",
               date_fin_validite as "Fin de validité", statut_classement as "Statut"
        from marts.fct_hebergements_classes
        where statut_classement in ('Expiré', 'Expire sous 6 mois')
        order by date_fin_validite
    """)
    anomalies = load.lire("""
        select severite as "Sévérité", controle as "Contrôle", nom_commercial as "Établissement",
               commune_source as "Commune", detail as "Détail"
        from marts.mart_qualite_anomalies
        order by case severite when 'Bloquante' then 1 when 'Avertissement' then 2 else 3 end, controle
    """)

    with pd.ExcelWriter(chemin, engine="openpyxl") as writer:
        pd.DataFrame({"Note de synthèse": synthese.splitlines()}).to_excel(
            writer, sheet_name="Synthèse", index=False)
        communes.to_excel(writer, sheet_name="Communes", index=False)
        epci.to_excel(writer, sheet_name="Intercommunalités", index=False)
        a_renouveler.to_excel(writer, sheet_name="Classements à renouveler", index=False)
        anomalies.to_excel(writer, sheet_name="Qualité des données", index=False)

    wb = load_workbook(chemin)
    for ws in wb.worksheets:
        _mettre_en_forme(ws, largeur_max=110 if ws.title == "Synthèse" else 60)
    ws = wb["Synthèse"]
    for row in ws.iter_rows(min_row=2):
        texte = row[0].value or ""
        if texte.startswith("#"):
            row[0].value = texte.lstrip("# ")
            row[0].font = Font(bold=True, size=14 if texte.startswith("# ") else 12, color=BLEU)
        else:
            row[0].value = texte.replace("**", "").strip("*")
    # Échelle de couleur sur le taux de fonction touristique
    ws = wb["Communes"]
    if ws.max_row > 1:
        ws.conditional_formatting.add(
            f"I2:I{ws.max_row}",
            ColorScaleRule(start_type="min", start_color="FFFFFF", end_type="max", end_color="5B9BD5"))
        for row in ws.iter_rows(min_row=2, min_col=3, max_col=8):
            for c in row:
                c.number_format = "# ##0"
    ws = wb["Qualité des données"]
    couleurs = {"Bloquante": "F8CBAD", "Avertissement": "FFE699", "Info": "DDEBF7"}
    for row in ws.iter_rows(min_row=2):
        if row[0].value in couleurs:
            row[0].fill = PatternFill("solid", fgColor=couleurs[row[0].value])
    wb.save(chemin)
    log.info("Rapport Excel généré : %s", chemin.name)
    return str(chemin)
